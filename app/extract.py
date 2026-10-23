"""Step 1 of the pipeline: a lecture file -> a list of slides with text.

Every piece of text keeps its slide number. That number is what the app
shows the student later ("read again: slide 12"), so never lose it.
"""

from pathlib import Path

import pymupdf
from pptx import Presentation

Slide = dict  # {"slide": int, "text": str}


def read_pdf(path: str) -> list[Slide]:
    slides: list[Slide] = []
    with pymupdf.open(path) as doc:
        for number, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                slides.append({"slide": number, "text": text})
    return slides


def read_pptx(path: str) -> list[Slide]:
    slides: list[Slide] = []
    for number, slide in enumerate(Presentation(path).slides, start=1):
        parts = [
            shape.text.strip()
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        ]
        if parts:
            slides.append({"slide": number, "text": "\n".join(parts)})
    return slides


def read_lecture(path: str) -> list[Slide]:
    """Read a .pdf or .pptx lecture. Raises if the file type is not supported."""
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        return read_pdf(path)
    if suffix == ".pptx":
        return read_pptx(path)
    raise ValueError(f"Cannot read {suffix} files yet. Use .pdf or .pptx.")


def split_for_prompt(slides: list[Slide], max_chars: int = 30_000) -> list[list[Slide]]:
    """Cut a long lecture into parts that each fit into one request to the model."""
    parts: list[list[Slide]] = [[]]
    used = 0
    for slide in slides:
        size = len(slide["text"])
        if parts[-1] and used + size > max_chars:
            parts.append([])
            used = 0
        parts[-1].append(slide)
        used += size
    return parts


def as_prompt_text(slides: list[Slide], max_chars: int = 40_000) -> str:
    """Join the slides into one text for the model, keeping slide numbers.

    max_chars protects you from sending a 300-page book to the API by mistake.
    When one lecture stops fitting here, that is the moment to add
    vector search (pgvector) instead of sending everything.
    """
    out: list[str] = []
    used = 0
    for slide in slides:
        block = f"[slide {slide['slide']}]\n{slide['text']}"
        if used + len(block) > max_chars:
            break
        out.append(block)
        used += len(block)
    return "\n\n".join(out)
